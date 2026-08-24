#!/usr/bin/env python3
"""Fill the repository-owned sanitized blank reference-check template."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from typing import Any, Iterable
from xml.sax.saxutils import escape

try:
    from docx import Document
except ImportError:  # The bundled runtime can use the deterministic OOXML fallback.
    Document = None


MISSING_ANSWER = "Not discussed during the reference check."
TEMPLATE_NAME = "reference-check-template.docx"

ANSWER_KEYS = [
    "known_duration",
    "working_capacity",
    "overall_performance",
    "responsibilities",
    "area_for_improvement",
    "performance_rating",
    "communication",
    "interactions",
    "teamwork",
    "adaptability",
    "problem_solving",
    "dependability",
    "leadership_potential",
    "recommendation",
    "rehire",
    "additional_comments",
]


def required(value: Any, label: str) -> str:
    result = "" if value is None else str(value).strip()
    if not result:
        raise ValueError(f"Missing required field: {label}")
    return result


def clean(value: Any) -> str:
    result = "" if value is None else str(value).strip()
    return result or MISSING_ANSWER


def require_object(data: dict[str, Any], key: str) -> dict[str, Any]:
    value = data.get(key, {})
    if not isinstance(value, dict):
        raise ValueError(f"'{key}' must be a JSON object")
    return value


def paragraphs(document: Any) -> Iterable[Any]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_token(document: Any, token: str, value: str) -> None:
    matches = []
    for paragraph in paragraphs(document):
        for run in paragraph.runs:
            if token in run.text:
                matches.append(run)
    if len(matches) != 1:
        raise ValueError(f"Sanitized template token {token} expected once; found {len(matches)}")
    matches[0].text = matches[0].text.replace(token, value)


def assert_sanitized_template(document: Any) -> None:
    text = "\n".join(paragraph.text for paragraph in paragraphs(document))
    required_tokens = {
        "{{candidate_name}}",
        "{{reference_name}}",
        "{{professional_relationship}}",
        "{{recommendation}}",
        "{{completed_by}}",
        "{{date}}",
    }
    missing = sorted(token for token in required_tokens if text.count(token) != 1)
    if missing:
        raise ValueError(f"Sanitized template contract failed for: {', '.join(missing)}")
    if "Professional Reference Check Form" not in text:
        raise ValueError("Sanitized template title is missing")


def keep_questions_with_answers(document: Any) -> None:
    for paragraph in paragraphs(document):
        if paragraph.text.strip().endswith("?"):
            paragraph.paragraph_format.keep_with_next = True


def keep_questions_with_answers_ooxml(document_xml: str) -> str:
    def update(match: re.Match[str]) -> str:
        paragraph_xml = match.group(0)
        text = "".join(re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", paragraph_xml, flags=re.DOTALL))
        text = re.sub(r"<[^>]+>", "", text).strip()
        if not text.endswith("?") or "<w:keepNext" in paragraph_xml:
            return paragraph_xml
        if "<w:pPr" in paragraph_xml:
            return re.sub(r"(<w:pPr(?:\s[^>]*)?>)", r"\1<w:keepNext/>", paragraph_xml, count=1)
        return re.sub(r"(<w:p(?:\s[^>]*)?>)", r"\1<w:pPr><w:keepNext/></w:pPr>", paragraph_xml, count=1)

    return re.sub(r"<w:p(?:\s[^>]*)?>.*?</w:p>", update, document_xml, flags=re.DOTALL)


def build_values(data: dict[str, Any]) -> dict[str, str]:
    candidate = require_object(data, "candidate")
    reference = require_object(data, "reference")
    answers = require_object(data, "answers")
    values = {
        "candidate_name": required(candidate.get("full_name"), "candidate.full_name"),
        "position_applied_for": required(candidate.get("position_applied_for"), "candidate.position_applied_for"),
        "client_company": required(candidate.get("company_name"), "candidate.company_name"),
        "reference_name": required(reference.get("full_name"), "reference.full_name"),
        "reference_title": required(reference.get("job_title"), "reference.job_title"),
        "reference_company": required(reference.get("company_name"), "reference.company_name"),
        "professional_relationship": required(reference.get("professional_relationship"), "reference.professional_relationship"),
        "completed_by": required(data.get("completed_by"), "completed_by"),
        "date": required(data.get("date"), "date"),
        **{key: clean(answers.get(key)) for key in ANSWER_KEYS},
    }
    strengths = answers.get("strengths")
    if strengths is None:
        strengths = []
    if not isinstance(strengths, list):
        strengths = [strengths]
    strengths = [str(item).strip() for item in strengths if str(item).strip()][:5]
    if not strengths:
        strengths = [MISSING_ANSWER]
    for index in range(1, 6):
        values[f"strength_{index}"] = strengths[index - 1] if index <= len(strengths) else "Not provided."
    return values


def build_with_ooxml(template: Path, output: Path, values: dict[str, str]) -> None:
    """Replace one-token-per-run placeholders without requiring python-docx.

    The generated sanitized template deliberately keeps each token in one run.
    This fallback verifies that contract before making a byte-preserving DOCX copy.
    """
    with zipfile.ZipFile(template) as source:
        document_xml = source.read("word/document.xml").decode("utf-8")
        core_xml = source.read("docProps/core.xml").decode("utf-8")
        required_tokens = {"{{candidate_name}}", "{{reference_name}}", "{{professional_relationship}}", "{{recommendation}}", "{{completed_by}}", "{{date}}"}
        missing = sorted(token for token in required_tokens if document_xml.count(token) != 1)
        if missing or "Professional Reference Check Form" not in document_xml:
            raise ValueError(f"Sanitized template contract failed for: {', '.join(missing) or 'title'}")
        for key, value in values.items():
            token = f"{{{{{key}}}}}"
            if document_xml.count(token) != 1:
                raise ValueError(f"Sanitized template token {token} expected once; found {document_xml.count(token)}")
            document_xml = document_xml.replace(token, escape(value))
        document_xml = keep_questions_with_answers_ooxml(document_xml)
        title = escape(f"{values['candidate_name']} Reference Check")
        core_xml = re.sub(r"(<dc:title>).*?(</dc:title>)", rf"\\g<1>{title}\\g<2>", core_xml, count=1)
        output.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "word/document.xml":
                    content = document_xml.encode("utf-8")
                elif item.filename == "docProps/core.xml":
                    content = core_xml.encode("utf-8")
                destination.writestr(item, content)


def build(data: dict[str, Any], output: Path, template: Path) -> None:
    if not template.exists():
        raise FileNotFoundError(f"Reference template not found: {template}")
    values = build_values(data)
    if Document is None:
        build_with_ooxml(template, output, values)
        return
    document = Document(template)
    assert_sanitized_template(document)

    for key, value in values.items():
        replace_token(document, f"{{{{{key}}}}}", value)
    keep_questions_with_answers(document)

    core = document.core_properties
    core.title = f"{values['candidate_name']} Reference Check"
    core.subject = "Professional candidate reference check"
    core.author = "Top Tier Talent Group"
    core.keywords = "reference check, candidate, recruitment"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="JSON input file")
    parser.add_argument("--output", required=True, type=Path, help="DOCX output file")
    parser.add_argument("--template", type=Path, help="Optional sanitized template path")
    args = parser.parse_args()
    if args.output.suffix.lower() != ".docx":
        parser.error("--output must end in .docx")
    data = json.loads(args.input.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Input JSON must contain one top-level object")
    template = args.template or Path(__file__).resolve().parent.parent / "assets" / TEMPLATE_NAME
    build(data, args.output.resolve(), template.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
